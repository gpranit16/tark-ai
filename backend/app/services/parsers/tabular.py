import csv
import io
from uuid import UUID

import openpyxl

from app.schemas.document import Block, Page, TableData, ParsedDocumentData
from app.services.parsers.base import BaseParser


class TabularParser(BaseParser):
    async def parse(
        self, content_bytes: bytes, file_id: UUID, filename: str, extension: str
    ) -> tuple[ParsedDocumentData, bool]:
        ext = extension.lower()
        pages: list[Page] = []

        if ext == ".csv":
            try:
                text_content = content_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text_content = content_bytes.decode("latin-1", errors="replace")

            reader = csv.reader(io.StringIO(text_content))
            all_rows = [row for row in reader if row]
            
            headers = all_rows[0] if all_rows else []
            data_rows = all_rows[1:] if len(all_rows) > 1 else []

            table = TableData(page_number=1, headers=headers, rows=data_rows)
            table_summary = f"CSV Table ({len(all_rows)} rows, {len(headers)} columns)"
            block = Block(
                block_type="table",
                content=table_summary,
                page_number=1,
                metadata={"headers": headers, "row_count": len(data_rows)}
            )

            raw_text = "\n".join([", ".join(r) for r in all_rows])
            page = Page(
                page_number=1,
                text=raw_text,
                blocks=[block],
                tables=[table],
                images=[]
            )
            pages.append(page)

        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
            for page_num, sheet_name in enumerate(wb.sheetnames, start=1):
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                clean_rows = [[str(val) if val is not None else "" for val in r] for r in rows if any(val is not None for val in r)]
                
                if not clean_rows:
                    continue

                headers = clean_rows[0]
                data_rows = clean_rows[1:] if len(clean_rows) > 1 else []

                table = TableData(page_number=page_num, headers=headers, rows=data_rows)
                block = Block(
                    block_type="table",
                    content=f"Sheet '{sheet_name}' ({len(clean_rows)} rows)",
                    page_number=page_num,
                    metadata={"sheet_name": sheet_name, "headers": headers}
                )

                sheet_text = "\n".join([", ".join(r) for r in clean_rows])
                pages.append(Page(
                    page_number=page_num,
                    text=sheet_text,
                    blocks=[block],
                    tables=[table],
                    images=[]
                ))
            
            wb.close()

        if not pages:
            pages.append(Page(page_number=1, text="", blocks=[], tables=[], images=[]))

        doc_data = ParsedDocumentData(
            file_id=file_id,
            filename=filename,
            extension=extension,
            pages=pages,
            metadata={"parser": "TabularParser", "total_sheets_or_pages": len(pages)}
        )
        return doc_data, False
