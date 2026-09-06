import io
from uuid import UUID
from pptx import Presentation

from app.schemas.document import Block, Page, TableData, ParsedDocumentData
from app.services.parsers.base import BaseParser


class PptxParser(BaseParser):
    async def parse(
        self, content_bytes: bytes, file_id: UUID, filename: str, extension: str
    ) -> tuple[ParsedDocumentData, bool]:
        prs = Presentation(io.BytesIO(content_bytes))
        pages: list[Page] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_blocks: list[Block] = []
            slide_tables: list[TableData] = []
            slide_text_parts: list[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        slide_blocks.append(Block(
                            block_type="paragraph",
                            content=text,
                            page_number=slide_idx
                        ))
                        slide_text_parts.append(text)

                if shape.has_table:
                    tbl = shape.table
                    rows: list[list[str]] = []
                    for r in tbl.rows:
                        row_cells = [cell.text.strip() for cell in r.cells]
                        rows.append(row_cells)
                    
                    if rows:
                        headers = rows[0]
                        data_rows = rows[1:] if len(rows) > 1 else []
                        tbl_data = TableData(page_number=slide_idx, headers=headers, rows=data_rows)
                        slide_tables.append(tbl_data)
                        
                        tbl_summary = f"Table ({len(rows)} rows)"
                        slide_blocks.append(Block(
                            block_type="table",
                            content=tbl_summary,
                            page_number=slide_idx,
                            metadata={"headers": headers}
                        ))
                        slide_text_parts.append(tbl_summary)

            slide_text = "\n".join(slide_text_parts)
            pages.append(Page(
                page_number=slide_idx,
                text=slide_text,
                blocks=slide_blocks,
                tables=slide_tables,
                images=[]
            ))

        if not pages:
            pages.append(Page(page_number=1, text="", blocks=[], tables=[], images=[]))

        doc_data = ParsedDocumentData(
            file_id=file_id,
            filename=filename,
            extension=extension,
            pages=pages,
            metadata={"parser": "PptxParser", "total_slides": len(prs.slides)}
        )
        return doc_data, False
