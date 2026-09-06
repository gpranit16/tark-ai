import io
from uuid import UUID
import docx

from app.schemas.document import Block, Page, TableData, ParsedDocumentData
from app.services.parsers.base import BaseParser


class DocxParser(BaseParser):
    async def parse(
        self, content_bytes: bytes, file_id: UUID, filename: str, extension: str
    ) -> tuple[ParsedDocumentData, bool]:
        doc = docx.Document(io.BytesIO(content_bytes))
        blocks: list[Block] = []
        tables: list[TableData] = []
        full_text_lines: list[str] = []

        # Parse paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else ""
            b_type = "heading" if "heading" in style_name.lower() else "paragraph"

            blocks.append(Block(
                block_type=b_type,
                content=text,
                page_number=1,
                metadata={"style": style_name}
            ))
            full_text_lines.append(text)

        # Parse tables
        for tbl in doc.tables:
            table_rows: list[list[str]] = []
            for row in tbl.rows:
                cell_texts = [c.text.strip() for c in row.cells]
                table_rows.append(cell_texts)

            if table_rows:
                headers = table_rows[0]
                data_rows = table_rows[1:] if len(table_rows) > 1 else []
                table_data = TableData(page_number=1, headers=headers, rows=data_rows)
                tables.append(table_data)

                table_summary = f"Table ({len(table_rows)} rows)"
                blocks.append(Block(
                    block_type="table",
                    content=table_summary,
                    page_number=1,
                    metadata={"headers": headers}
                ))
                full_text_lines.append(table_summary)

        full_text = "\n".join(full_text_lines)
        page = Page(
            page_number=1,
            text=full_text,
            blocks=blocks,
            tables=tables,
            images=[]
        )

        doc_data = ParsedDocumentData(
            file_id=file_id,
            filename=filename,
            extension=extension,
            pages=[page],
            metadata={"parser": "DocxParser", "total_paragraphs": len(doc.paragraphs), "total_tables": len(doc.tables)}
        )
        return doc_data, False
