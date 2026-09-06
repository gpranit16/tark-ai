import io
from uuid import uuid4

import pytest
import pypdf
import docx
from pptx import Presentation
import openpyxl
from PIL import Image, ImageDraw
from fastapi.testclient import TestClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.document import ParsedDocument
from app.services.parsers.factory import get_parser_for_extension


@pytest.fixture
def mock_storage_and_ocr(monkeypatch):
    storage_dict = {}

    class MockStorage:
        @property
        def provider_name(self): return "local"
        
        async def upload(self, file_stream, key, mime):
            storage_dict[key] = file_stream.read()

        async def download(self, key):
            if key in storage_dict:
                return storage_dict[key]
            return b"fake stream data"

        async def delete(self, key):
            storage_dict.pop(key, None)

        async def exists(self, key):
            return key in storage_dict

    class MockOCR:
        @property
        def provider_name(self): return "mock_ocr"

        async def extract_text(self, image_bytes: bytes):
            return "Sample OCR Extracted Text", [{"text": "Sample OCR Extracted Text", "confidence": 0.99}]

    monkeypatch.setattr("app.services.files.get_storage_provider", lambda *args, **kwargs: MockStorage())
    monkeypatch.setattr("app.api.v1.files.get_storage_provider", lambda *args, **kwargs: MockStorage())
    monkeypatch.setattr("app.services.document_processing.get_storage_provider", lambda *args, **kwargs: MockStorage())
    monkeypatch.setattr("app.services.ocr.factory.get_ocr_provider", lambda: MockOCR())
    monkeypatch.setattr("app.services.parsers.image.get_ocr_provider", lambda: MockOCR())
    monkeypatch.setattr("app.services.parsers.pdf.get_ocr_provider", lambda: MockOCR())


# 1. Test Text Parser (TXT & MD)
@pytest.mark.asyncio
async def test_text_parser():
    parser = get_parser_for_extension(".txt")
    content = b"# Header Title\n\nThis is a sample paragraph.\n\nAnother line."
    file_id = uuid4()
    doc_data, ocr_used = await parser.parse(content, file_id, "test.txt", ".txt")
    
    assert doc_data.file_id == file_id
    assert len(doc_data.pages) == 1
    assert doc_data.pages[0].blocks[0].block_type == "heading"
    assert doc_data.pages[0].blocks[0].content == "# Header Title"
    assert ocr_used is False


# 2. Test Tabular Parser (CSV & XLSX)
@pytest.mark.asyncio
async def test_csv_parser():
    parser = get_parser_for_extension(".csv")
    csv_bytes = b"Name,Age,Role\nAlice,30,Engineer\nBob,25,Designer"
    file_id = uuid4()
    doc_data, ocr_used = await parser.parse(csv_bytes, file_id, "data.csv", ".csv")

    assert len(doc_data.pages) == 1
    table = doc_data.pages[0].tables[0]
    assert table.headers == ["Name", "Age", "Role"]
    assert len(table.rows) == 2
    assert ocr_used is False


@pytest.mark.asyncio
async def test_xlsx_parser():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Item", "Price"])
    ws.append([1, "Widget", 19.99])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    xlsx_bytes = buf.read()

    parser = get_parser_for_extension(".xlsx")
    file_id = uuid4()
    doc_data, _ = await parser.parse(xlsx_bytes, file_id, "data.xlsx", ".xlsx")

    assert len(doc_data.pages) == 1
    assert doc_data.pages[0].tables[0].headers == ["ID", "Item", "Price"]


# 3. Test DOCX Parser
@pytest.mark.asyncio
async def test_docx_parser():
    doc = docx.Document()
    doc.add_heading("Project Report", level=1)
    doc.add_paragraph("This is the main introduction paragraph.")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Header A"
    tbl.cell(0, 1).text = "Header B"
    tbl.cell(1, 0).text = "Val A"
    tbl.cell(1, 1).text = "Val B"
    
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    parser = get_parser_for_extension(".docx")
    file_id = uuid4()
    doc_data, _ = await parser.parse(docx_bytes, file_id, "doc.docx", ".docx")

    assert len(doc_data.pages) == 1
    assert len(doc_data.pages[0].blocks) >= 2
    assert len(doc_data.pages[0].tables) == 1
    assert doc_data.pages[0].tables[0].headers == ["Header A", "Header B"]


# 4. Test PPTX Parser
@pytest.mark.asyncio
async def test_pptx_parser():
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    txBox = slide.shapes.add_textbox(0, 0, 100, 100)
    txBox.text_frame.text = "Presentation Slide 1 Title"

    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    parser = get_parser_for_extension(".pptx")
    file_id = uuid4()
    doc_data, _ = await parser.parse(pptx_bytes, file_id, "pres.pptx", ".pptx")

    assert len(doc_data.pages) == 1
    assert "Presentation Slide 1 Title" in doc_data.pages[0].text


# 5. Test Normal PDF Parser
@pytest.mark.asyncio
async def test_pdf_parser_normal():
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()

    parser = get_parser_for_extension(".pdf")
    file_id = uuid4()
    doc_data, ocr_used = await parser.parse(pdf_bytes, file_id, "sample.pdf", ".pdf")

    assert doc_data.file_id == file_id
    assert len(doc_data.pages) == 1


# 6. Test Image OCR Parser
@pytest.mark.asyncio
async def test_image_ocr_parser(mock_storage_and_ocr):
    img = Image.new("RGB", (100, 50), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((10, 10), "TEST OCR", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    img_bytes = buf.getvalue()

    parser = get_parser_for_extension(".png")
    file_id = uuid4()
    doc_data, ocr_used = await parser.parse(img_bytes, file_id, "image.png", ".png")

    assert ocr_used is True
    assert "Sample OCR Extracted Text" in doc_data.pages[0].text
    assert len(doc_data.pages[0].images) == 1


# 7. End-to-End File Upload and Processing Flow
@pytest.mark.asyncio
async def test_end_to_end_upload_and_processing(client: TestClient, user, mock_storage_and_ocr, db_session: AsyncSession):
    content = b"# End to End Test Document\n\nThis is a complete async document parsing test."
    files = {"file": ("e2e_doc.md", io.BytesIO(content), "text/markdown")}
    data = {"user_id": str(user.id)}

    # Upload file
    response = client.post("/api/v1/files/upload", files=files, data=data)
    assert response.status_code == 201
    file_id = response.json()["id"]

    # Check status endpoint
    status_resp = client.get(f"/api/v1/documents/{file_id}/status?user_id={user.id}")
    assert status_resp.status_code == 200
    assert status_resp.json()["status"] in ("pending", "processing", "completed")

    # Fetch parsed document content
    doc_resp = client.get(f"/api/v1/documents/{file_id}?user_id={user.id}")
    assert doc_resp.status_code == 200
    doc_json = doc_resp.json()
    assert doc_json["file_id"] == file_id

    # Fetch metadata
    meta_resp = client.get(f"/api/v1/documents/{file_id}/metadata?user_id={user.id}")
    assert meta_resp.status_code == 200
    assert meta_resp.json()["filename"] == "e2e_doc.md"


# 8. Test Ownership Isolation on Document Endpoints
@pytest.mark.asyncio
async def test_document_ownership_isolation(client: TestClient, user, mock_storage_and_ocr):
    content = b"User A private text."
    files = {"file": ("private.txt", io.BytesIO(content), "text/plain")}
    data = {"user_id": str(user.id)}
    upload_resp = client.post("/api/v1/files/upload", files=files, data=data)
    file_id = upload_resp.json()["id"]

    wrong_user = str(uuid4())

    # Unauthorized access returns 404
    assert client.get(f"/api/v1/documents/{file_id}/status?user_id={wrong_user}").status_code == 404
    assert client.get(f"/api/v1/documents/{file_id}?user_id={wrong_user}").status_code == 404
    assert client.get(f"/api/v1/documents/{file_id}/metadata?user_id={wrong_user}").status_code == 404
    assert client.post(f"/api/v1/documents/{file_id}/retry?user_id={wrong_user}").status_code == 404


# 9. Test Retry Flow
@pytest.mark.asyncio
async def test_retry_flow(client: TestClient, user, mock_storage_and_ocr, db_session: AsyncSession):
    content = b"Sample document text for retry flow."
    files = {"file": ("retry_test.txt", io.BytesIO(content), "text/plain")}
    data = {"user_id": str(user.id)}
    upload_resp = client.post("/api/v1/files/upload", files=files, data=data)
    file_id = upload_resp.json()["id"]

    # Trigger retry
    retry_resp = client.post(f"/api/v1/documents/{file_id}/retry?user_id={user.id}")
    assert retry_resp.status_code == 200
    assert retry_resp.json()["status"] in ("pending", "processing", "completed")
